"""Phase two: timed rewards, balance tiers and audited website management."""
import csv
import io
import re
from datetime import datetime
from uuid import uuid4
from zoneinfo import ZoneInfo

from .settings import CAMPAIGN_TIMEZONE, MAX_TOTAL_TOKENS_LIMIT, MINING_DEFAULT_BUDGET, MINING_SECONDS, REQUIRED_REFERRALS, TIERS

BONUS_BATCH_SIZE = 200
# A finished mining cycle must be claimed within this window. After it, the
# reserved reward returns to the mining budget instead of staying locked forever.
MINING_CLAIM_WINDOW = 7 * 86400


def level_for(user):
    from .service import earned
    balance = max(0, earned(user) - max(0, int(user.get("paid_amount", 0))))
    refs = max(0, int(user.get("ref_count", 0)))
    index = max(i for i, tier in enumerate(TIERS) if refs >= tier["referrals"] and balance >= tier["balance"])
    return {**TIERS[index], "stored_balance": balance, "next": dict(TIERS[index + 1]) if index + 1 < len(TIERS) else None}


class PhaseTwoMixin:
    def admin_bot_bonus(self, admin_id, operation_id, amount, user_id=None):
        """Legacy Telegram controls use the same transactional reward budget."""
        from .service import PortalError
        self.require_admin(admin_id)
        if type(amount) is not int or not 1 <= amount <= MAX_TOTAL_TOKENS_LIMIT or not isinstance(operation_id, str) or not re.fullmatch(r"[a-f0-9]{12,64}", operation_id):
            raise PortalError("invalid_request")
        receipt = f"bot-bonus:{admin_id}:{operation_id}"
        field = "manual_bonus_tokens" if user_id is not None else "bonus_tokens"
        ids_field = "manual_bonus_operation_ids" if user_id is not None else "bonus_operation_ids"
        history_field = "manual_bonus_history" if user_id is not None else "bonus_history"

        def eligible(u):
            return (int(u.get("ref_count", 0) or 0) >= REQUIRED_REFERRALS and int(u.get("submitted", 0) or 0) > 0
                    and not u.get("portal_suspended") and re.fullmatch(r"0x[a-fA-F0-9]{40}", u.get("wallet") or "")
                    and operation_id not in (u.get(ids_field) or []))

        def credit(tx, users):
            self._credit_lock(tx, amount * len(users))
            for user in users:
                user[field] = int(user.get(field, 0)) + amount
                user["paid"] = 0
                user.setdefault(ids_field, []).append(operation_id)
                user.setdefault(history_field, []).append({"operation_id": operation_id, "amount": amount,
                    "created_at": self.now(), "created_by": admin_id})
                tx.put("users", user["user_id"], user)
                self._history(tx, f"{receipt}:{user['user_id']}", user["user_id"], "admin_bonus", amount,
                              "پاداش مدیریت · Management reward")

        def finish(tx, count):
            result = {"user_id": user_id, "amount": amount, "count": count, "total": amount * count}
            tx.put("portal_admin_operations", receipt, result)
            self._audit(tx, admin_id, "bonus", user_id if user_id is not None else "campaign",
                        {"amount": amount, "count": count, "source": "telegram"})
            return result

        def previous_result(tx):
            previous = tx.get("portal_admin_operations", receipt)
            if previous and (previous["amount"] != amount or previous["user_id"] != user_id):
                raise PortalError("operation_conflict", 409)
            return previous

        if user_id is not None:
            def single(tx):
                previous = previous_result(tx)
                if previous:
                    return {**previous, "already_completed": True}
                user = tx.get("users", user_id)
                if not user:
                    raise PortalError("user_not_found", 404)
                if user.get("portal_suspended"):
                    raise PortalError("account_suspended", 403)
                users = [user] if operation_id not in user.get(ids_field, []) else []
                credit(tx, users)
                return {**finish(tx, len(users)), "already_completed": not users}
            return self.store.transaction(single)

        # Campaign-wide bonus: one huge transaction fails on a large user base,
        # so credit in batches. A rerun with the same operation_id skips users
        # already credited, so an interrupted run can be resumed safely.
        previous = self.store.transaction(previous_result)
        if previous:
            return {**previous, "already_completed": True}
        candidates = [u["user_id"] for u in self.store.scan(
            "users", ("user_id", "ref_count", "submitted", "portal_suspended", "wallet", ids_field)) if eligible(u)]
        # Refuse up front when the whole operation cannot fit in the campaign cap.
        self.store.transaction(lambda tx: self._credit_lock(tx, amount * len(candidates), commit=False))
        credited = 0
        for start in range(0, len(candidates), BONUS_BATCH_SIZE):
            batch = candidates[start:start + BONUS_BATCH_SIZE]
            def batch_operation(tx, batch=batch):
                users = [u for u in (tx.get("users", uid) for uid in batch) if u and eligible(u)]
                credit(tx, users)
                return len(users)
            credited += self.store.transaction(batch_operation)
        return {**self.store.transaction(lambda tx: finish(tx, credited)), "already_completed": False}

    def admin_bot_payment(self, admin_id, user_id, action, paid_total=None):
        """Record the trusted admin's legacy payment status without write skew.

        This records an external payment/status correction, never sends tokens.
        A pending portal request must be reviewed through its existing workflow.
        paid_total (from a payment file) marks only that much as paid, so
        rewards earned after the file was prepared stay unpaid.
        """
        from .service import PortalError, earned
        self.require_admin(admin_id)
        if action not in ("yes", "no", "unpaid"):
            raise PortalError("invalid_request")
        if paid_total is not None and (type(paid_total) is not int or paid_total < 0):
            raise PortalError("invalid_request")
        receipt = f"bot-payment:{uuid4().hex}"
        def operation(tx):
            user = tx.get("users", user_id)
            if not user:
                raise PortalError("user_not_found", 404)
            if tx.get("portal_open_payments", user_id):
                raise PortalError("payment_pending", 409)
            before = max(0, int(user.get("paid_amount", 0)))
            if action == "yes":
                total = earned(user)
                target = total if paid_total is None else min(total, max(before, paid_total))
                user.update(paid=int(target >= total), paid_amount=target)
            elif action == "no":
                user["paid"] = 0
            else:
                # Tokens confirmed on-chain cannot be "unpaid": resetting would let
                # the user withdraw the same balance a second time.
                if tx.find("portal_payments", {"user_id": user_id, "status": "recorded"}, limit=1):
                    raise PortalError("onchain_payment_recorded", 409)
                user.update(paid=0, paid_amount=0)
            after = max(0, int(user.get("paid_amount", 0)))
            tx.put("users", user_id, user)
            if before != after:
                self._history(tx, receipt, user_id, "payment", before-after,
                              "ثبت وضعیت پرداخت توسط مدیر · Admin payment status")
            self._audit(tx, admin_id, "review_payment", user_id, {"source": "telegram", "action": action,
                        "paid_before": before, "paid_after": after})
            return {"paid": after}
        return self.store.transaction(operation)

    def register_gateway_user(self, user_id, referrer_id):
        """Complete a solved captcha + membership flow without bot notifications.

        The user document and inviter count commit together. Concurrent replies
        or old join buttons cannot credit the same invitation twice.
        """
        from .service import PortalError, earned
        def operation(tx):
            # Legacy users do not necessarily have a unique user_id index.
            # Serialize registration upserts even when there is no inviter.
            lock = tx.get("portal_state", "registration_lock") or {"sequence": 0}
            lock["sequence"] += 1
            tx.put("portal_state", "registration_lock", lock)
            user = tx.get("users", user_id) or {"user_id": user_id, "ref_count": 0, "lang": "fa"}
            if user.get("portal_suspended"):
                raise PortalError("account_suspended", 403)
            candidate = user.get("registration_referrer") or referrer_id
            if type(candidate) is int and candidate > 0 and candidate != user_id and not user.get("referred_by") and int(user.get("verified", 0)) < 1:
                inviter = tx.get("users", candidate) or {"user_id": candidate, "verified": 0, "ref_count": 0}
                old_total = earned(inviter)
                inviter["ref_count"] = int(inviter.get("ref_count", 0)) + 1
                delta = earned(inviter) - old_total
                self._credit_lock(tx, delta)
                tx.put("users", candidate, inviter)
                user["referred_by"] = candidate
                self._history(tx, f"referral:{user_id}:{candidate}", candidate, "referral", delta,
                              "دعوت تأییدشده · Verified referral")
            user["verified"] = max(1, int(user.get("verified", 0)))
            user.setdefault("registered_at", self.now())
            tx.put("users", user_id, user)
        self.store.transaction(operation)

    def check_access(self, user_id):
        from .service import PortalError
        if (self.store.get("users", user_id) or {}).get("portal_suspended"):
            raise PortalError("account_suspended", 403)
        # Every authenticated web API request rechecks channel membership, using
        # a short cache so page views do not hit Telegram's rate limit.
        if user_id not in self.settings.admins:
            self._member(user_id, fresh=False)

    def mining_config(self, tx=None):
        tx = tx or self.store
        return {"enabled": True, "budget": MINING_DEFAULT_BUDGET, "spent": 0, "reserved": 0,
                **(tx.get("portal_state", "mining") or {})}

    def mining_status(self, user_id, user=None):
        user = user or self.store.get("users", user_id) or {"user_id": user_id}
        config = self.mining_config()
        cycle = self.store.get("portal_mining", user_id)
        active = cycle if cycle and cycle.get("status") == "active" and not self._mining_expired(cycle) else None
        return {"enabled": config["enabled"], "duration": MINING_SECONDS, "level": level_for(user),
                "tiers": [dict(tier) for tier in TIERS], "total_claimed": int(user.get("mining_total", 0)),
                "budget_remaining": max(0, config["budget"] - config["spent"] - config["reserved"]),
                "cycle": {k: active[k] for k in ("id", "started_at", "ends_at", "reward", "tier")} if active else None}

    def start_mining(self, user_id):
        from .service import PortalError
        self._member(user_id)
        cycle_id = uuid4().hex
        def operation(tx):
            self._enabled(tx)
            user = self._user(tx, user_id)
            if int(user.get("ref_count", 0)) < REQUIRED_REFERRALS:
                raise PortalError("referrals_required", 403)
            if not re.fullmatch(r"0x[a-fA-F0-9]{40}", user.get("wallet") or ""):
                raise PortalError("wallet_required", 403)
            current = tx.get("portal_mining", user_id)
            if current and current.get("status") == "active":
                if not self._mining_expired(current):
                    raise PortalError("mining_active", 409)
                self._release_mining_cycle(tx, current)
            config = self.mining_config(tx)
            if not config["enabled"]:
                raise PortalError("mining_paused", 409)
            tier = level_for(user)
            amount = tier["reward"]
            if config["spent"] + config["reserved"] + amount > config["budget"]:
                raise PortalError("mining_budget_exhausted", 409)
            # Reserve the reward before the timer starts. This only checks the
            # campaign cap; the reward is counted as earned when it is claimed.
            self._credit_lock(tx, amount, commit=False)
            config["reserved"] += amount
            tx.put("portal_state", "mining", config)
            cycle = {"id": cycle_id, "user_id": user_id, "status": "active", "started_at": self.now(),
                     "ends_at": self.now() + MINING_SECONDS, "reward": amount, "tier": tier["id"]}
            # Conflict with a concurrent completed payment or suspension so the
            # stored reward snapshots a consistent level and access state.
            user["portal_mining_version"] = int(user.get("portal_mining_version", 0)) + 1
            tx.put("users", user_id, user)
            tx.put("portal_mining", user_id, cycle)
            return {"cycle": cycle}
        return self.store.transaction(operation)

    def claim_mining(self, user_id, cycle_id):
        from .service import PortalError
        if not isinstance(cycle_id, str) or not re.fullmatch(r"[a-f0-9]{32}", cycle_id):
            raise PortalError("invalid_mining_cycle")
        self._member(user_id)
        def operation(tx):
            self._enabled(tx)
            user = self._user(tx, user_id)
            if int(user.get("ref_count", 0)) < REQUIRED_REFERRALS:
                raise PortalError("referrals_required", 403)
            if not re.fullmatch(r"0x[a-fA-F0-9]{40}", user.get("wallet") or ""):
                raise PortalError("wallet_required", 403)
            cycle = tx.get("portal_mining", user_id)
            receipt = f"mining:{user_id}:{cycle_id}"
            previous = tx.get("portal_history", receipt)
            if previous:
                return {"amount": previous["amount"], "already_claimed": True}
            if not cycle or cycle["id"] != cycle_id or cycle["status"] != "active":
                raise PortalError("mining_not_started", 409)
            if self.now() < cycle["ends_at"]:
                raise PortalError("mining_not_ready", 409, remaining=cycle["ends_at"] - self.now())
            if self._mining_expired(cycle):
                # Commit the release, then report the expiry outside the transaction.
                self._release_mining_cycle(tx, cycle)
                return {"_expired": True}
            amount = cycle["reward"]
            config = self.mining_config(tx)
            config["reserved"] -= amount
            config["spent"] += amount
            tx.put("portal_state", "mining", config)
            self._credit_lock(tx, amount)
            user["manual_bonus_tokens"] = int(user.get("manual_bonus_tokens", 0)) + amount
            user["mining_total"] = int(user.get("mining_total", 0)) + amount
            user["paid"] = 0
            tx.put("users", user_id, user)
            cycle.update(status="claimed", claimed_at=self.now())
            tx.put("portal_mining", user_id, cycle)
            self._history(tx, receipt, user_id, "mining", amount, "پاداش ماین · Mining reward")
            return {"amount": amount, "already_claimed": False}
        result = self.store.transaction(operation)
        if result.get("_expired"):
            raise PortalError("mining_expired", 409)
        return result

    def _mining_expired(self, cycle):
        return self.now() > int(cycle.get("ends_at", 0)) + MINING_CLAIM_WINDOW

    def _release_mining_cycle(self, tx, cycle):
        """Return an unclaimed, expired reservation to the mining budget."""
        config = self.mining_config(tx)
        config["reserved"] = max(0, int(config["reserved"]) - int(cycle.get("reward", 0)))
        tx.put("portal_state", "mining", config)
        cycle = dict(cycle, status="expired", expired_at=self.now())
        tx.put("portal_mining", cycle["user_id"], cycle)
        return cycle

    def expire_mining_cycles(self, limit=500):
        """Background sweep: release reservations of cycles nobody claimed in time."""
        cutoff = self.now() - MINING_CLAIM_WINDOW
        released = 0
        for cycle in self.store.find("portal_mining", {"status": "active", "ends_at": {"$lt": cutoff}}, limit=limit):
            def operation(tx, user_id=cycle["user_id"], cycle_id=cycle["id"]):
                current = tx.get("portal_mining", user_id)
                if current and current.get("id") == cycle_id and current.get("status") == "active" and self._mining_expired(current):
                    self._release_mining_cycle(tx, current)
                    return 1
                return 0
            released += self.store.transaction(operation)
        return released

    def _audit(self, tx, admin_id, action, target, details=None):
        key = uuid4().hex
        tx.put("portal_audit", key, {"id": key, "admin_id": admin_id, "action": action,
               "target": str(target), "details": details or {}, "created_at": self.now()})

    def campaign_controls(self):
        return {"paused": (self.store.get("bot_settings", "bot_status") or {}).get("status") == "off",
                "announcement": (self.store.get("portal_state", "announcement") or {}).get("text", ""),
                "mining": self.mining_config(), "withdrawals_enabled": (self.store.get("portal_state", "withdrawals") or {}).get("enabled", True), "timezone": CAMPAIGN_TIMEZONE}

    def admin_controls(self, admin_id, payload):
        from .service import PortalError
        self.require_admin(admin_id)
        allowed = {"paused", "mining_enabled", "mining_budget", "announcement", "withdrawals_enabled"}
        if not payload or not set(payload) <= allowed:
            raise PortalError("invalid_controls")
        for name in ("paused", "mining_enabled", "withdrawals_enabled"):
            if name in payload and type(payload[name]) is not bool:
                raise PortalError("invalid_controls")
        if "mining_budget" in payload and (type(payload["mining_budget"]) is not int or not 0 <= payload["mining_budget"] <= MAX_TOTAL_TOKENS_LIMIT):
            raise PortalError("invalid_controls")
        if "announcement" in payload and (not isinstance(payload["announcement"], str) or len(payload["announcement"]) > 1500):
            raise PortalError("invalid_controls")
        def operation(tx):
            if "paused" in payload:
                tx.put("bot_settings", "bot_status", {"key": "bot_status", "status": "off" if payload["paused"] else "on"})
            if "mining_enabled" in payload or "mining_budget" in payload:
                config = self.mining_config(tx)
                if "mining_enabled" in payload:
                    config["enabled"] = payload["mining_enabled"]
                if "mining_budget" in payload:
                    if payload["mining_budget"] < config["spent"] + config["reserved"]:
                        raise PortalError("budget_below_committed", 409)
                    config["budget"] = payload["mining_budget"]
                tx.put("portal_state", "mining", config)
            if "withdrawals_enabled" in payload:
                tx.put("portal_state", "withdrawals", {"enabled": payload["withdrawals_enabled"], "updated_at": self.now()})
            if "announcement" in payload:
                tx.put("portal_state", "announcement", {"text": payload["announcement"].strip(), "updated_at": self.now()})
            self._audit(tx, admin_id, "campaign_controls", "campaign", payload)
        self.store.transaction(operation)
        return self.campaign_controls()

    @staticmethod
    def _admin_user(user):
        from .service import MIN_WITHDRAWAL, earned
        wallet = user.get("wallet") or ""
        paid = int(user.get("paid_amount", 0))
        referrals = int(user.get("ref_count", 0))
        total_earned = earned(user)
        balance = max(0, total_earned - paid)
        wallet_valid = bool(re.fullmatch(r"0x[a-fA-F0-9]{40}", wallet))
        suspended = bool(user.get("portal_suspended"))
        eligible = (int(user.get("verified", 0)) >= 1 and referrals >= 3 and wallet_valid and
                    balance >= MIN_WITHDRAWAL and not suspended)
        return {"user_id": user["user_id"], "registered": int(user.get("verified", 0)) >= 1,
                "verified": int(user.get("verified", 0)),
                "telegram_username": user.get("telegram_username") or "",
                "lang": user.get("lang") or "",
                "registered_at": user.get("registered_at"),
                "referrals": referrals, "referred_by": user.get("referred_by"),
                "earned": total_earned, "paid": paid, "unpaid": balance,
                "bonus_tokens": int(user.get("bonus_tokens", 0)),
                "manual_bonus_tokens": int(user.get("manual_bonus_tokens", 0)),
                "daily_count": int(user.get("daily_count", 0)),
                "last_daily": user.get("last_daily", 0),
                "level": level_for(user), "wallet": wallet,
                "wallet_registered": bool(wallet), "submitted": int(user.get("submitted", 0)),
                "instagram": user.get("instagram") or user.get("instagram_id") or "",
                "twitter": user.get("twitter") or user.get("twitter_id") or "",
                "social_registered": bool(user.get("social_registered") or user.get("instagram_id") or user.get("twitter_id")),
                "social_registered_at": user.get("social_registered_at"),
                "social_updated_at": user.get("social_updated_at"),
                "suspended": suspended, "eligible": eligible,
                "mining_total": int(user.get("mining_total", 0)),
                "challenge_quiz_attempts": int(user.get("challenge_quiz_attempts", 0)),
                "challenge_quiz_passed": bool(user.get("challenge_quiz_passed", False))}

    def admin_users(self, admin_id, query="", page=1, eligible_only=False):
        from .service import MIN_WITHDRAWAL, PortalError
        self.require_admin(admin_id)
        if not isinstance(query, str) or len(query) > 100 or type(page) is not int or not 1 <= page <= 100000:
            raise PortalError("invalid_request")
        query = query.strip().lower()
        normalized_query = query.lstrip("@")
        users = []
        for raw_user in self.store.all("users"):
            user = self._admin_user(raw_user)
            if eligible_only and not user["eligible"]:
                continue
            if query and not any(normalized_query in str(value or "").lower() for value in (
                    user["user_id"], user["telegram_username"], user["wallet"],
                    user["instagram"], user["twitter"])):
                continue
            users.append(user)
        users.sort(key=lambda u: u["user_id"], reverse=True)
        total = len(users)
        return {"users": users[(page - 1) * 50:page * 50], "total": total,
                "page": page, "page_size": 50, "eligible_only": bool(eligible_only),
                "minimum_withdrawal": MIN_WITHDRAWAL if eligible_only else None}

    def admin_user_detail(self, admin_id, user_id):
        from .service import PortalError
        self.require_admin(admin_id)
        user = self.store.get("users", user_id)
        if not user:
            raise PortalError("user_not_found", 404)
        recent = lambda name: self.store.find(name, {"user_id": user_id}, sort=[("created_at", -1)], limit=100)
        invited = self.store.find("users", {"referred_by": user_id})
        referrals = [{"user_id": u["user_id"], "registered": int(u.get("verified", 0)) >= 1,
                      "telegram_username": u.get("telegram_username") or ""} for u in invited]
        # Referral quality: accounts that never claimed a daily reward and never
        # saved a wallet are a common sign of farmed (fake) invitations.
        inactive = sum(1 for u in invited if int(u.get("daily_count", 0) or 0) == 0 and not u.get("wallet"))
        detail_user = self._admin_user(user)
        detail_user["referral_quality"] = {"invited": len(invited), "inactive": inactive,
                                           "inactive_percent": round(100 * inactive / len(invited)) if invited else 0}
        return {"user": detail_user, "history": recent("portal_history"),
                "payments": recent("portal_payments"), "mining": self.mining_status(user_id, user),
                "submissions": recent("portal_submissions"), "referrals": referrals}
    def admin_user_action(self, admin_id, user_id, payload):
        from .service import PortalError
        self.require_admin(admin_id)
        action, reason = payload.get("action"), payload.get("reason", "")
        if action not in ("suspend", "resume", "unlock_wallet", "bonus"):
            raise PortalError("invalid_request")
        if not isinstance(reason, str) or not 3 <= len(reason.strip()) <= 400:
            raise PortalError("reason_required")
        amount, request_id = payload.get("amount"), payload.get("request_id")
        if action == "bonus" and (type(amount) is not int or not 1 <= amount <= MAX_TOTAL_TOKENS_LIMIT or
                                   not isinstance(request_id, str) or not re.fullmatch(r"[a-f0-9-]{20,64}", request_id)):
            raise PortalError("invalid_request")
        def operation(tx):
            user = tx.get("users", user_id)
            if not user:
                raise PortalError("user_not_found", 404)
            details = {"reason": reason.strip()}
            if action == "suspend":
                if user_id in self.settings.admins:
                    raise PortalError("admin_protected", 409)
                user["portal_suspended"] = True
            elif action == "resume":
                user["portal_suspended"] = False
            elif action == "unlock_wallet":
                if tx.get("portal_open_payments", user_id):
                    raise PortalError("wallet_locked", 409)
                user["submitted"] = 1 if user.get("wallet") else 0
            else:
                user = self._user(tx, user_id)
                receipt = f"admin-bonus:{admin_id}:{request_id}"
                existing = tx.get("portal_history", receipt)
                if existing:
                    if existing["user_id"] != user_id or existing["amount"] != amount:
                        raise PortalError("operation_conflict", 409)
                    return {"ok": True, "already_completed": True}
                self._credit_lock(tx, amount)
                user["manual_bonus_tokens"] = int(user.get("manual_bonus_tokens", 0)) + amount
                user["paid"] = 0
                self._history(tx, receipt, user_id, "admin_bonus", amount, reason.strip())
                details["amount"] = amount
            tx.put("users", user_id, user)
            self._audit(tx, admin_id, action, user_id, details)
            return {"ok": True}
        return self.store.transaction(operation)

    def export_report(self, admin_id, kind):
        from .service import PortalError
        self.require_admin(admin_id)
        if kind == "users":
            rows = [self._admin_user(u) for u in self.store.all("users")]
            fields = ["user_id", "registered", "referrals", "referred_by", "earned", "paid", "unpaid", "wallet", "suspended", "mining_total"]
        elif kind == "payments":
            rows = self.store.all("portal_payments")
            fields = ["id", "user_id", "amount", "wallet", "status", "tx_hash", "reason", "created_at", "reviewed_at", "reviewed_by"]
        elif kind == "audit":
            rows = [{**row, "details": str(row.get("details", {}))} for row in self.store.all("portal_audit")]
            fields = ["id", "admin_id", "action", "target", "details", "created_at"]
        else:
            raise PortalError("not_found", 404)
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(fields)
        for row in rows:
            values = []
            for field in fields:
                value = row.get(field, "")
                if field.endswith("_at") and value:
                    value = datetime.fromtimestamp(value, ZoneInfo(CAMPAIGN_TIMEZONE)).isoformat()
                value = str(value if value is not None else "")
                if value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
                    value = "'" + value
                values.append(value)
            writer.writerow(values)
        return "\ufeff" + output.getvalue()

    def export_xlsx(self, admin_id, kind):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from .service import PortalError
        self.require_admin(admin_id)
        if kind == "paid":
            rows = [p for p in self.store.all("portal_payments") if p.get("status") == "recorded"]
            fields = ["tx_hash", "user_id", "amount", "wallet", "status", "block_number", "created_at", "reviewed_at", "reviewed_by"]
            headers = ["Transaction hash", "Telegram numeric ID", "PRS paid", "Wallet", "Status", "Block", "Requested at", "Paid at", "Admin ID"]
        elif kind == "social":
            rows = [self._admin_user(u) for u in self.store.all("users")]
            fields = ["user_id", "instagram", "twitter"]
            headers = ["Telegram numeric ID", "Instagram URL", "X / Twitter URL"]
        else:
            raise PortalError("not_found", 404)
        book = Workbook()
        sheet = book.active
        sheet.title = "Paid transactions" if kind == "paid" else "Social profiles"
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="8B1E2D")
        for row in sorted(rows, key=lambda item: item.get("reviewed_at", item.get("user_id", 0)) or 0, reverse=True):
            values = []
            for field in fields:
                value = row.get(field, "")
                if field.endswith("_at") and value:
                    value = datetime.fromtimestamp(value, ZoneInfo(CAMPAIGN_TIMEZONE)).isoformat()
                values.append(value if value is not None else "")
            sheet.append(values)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(70, max(15, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
        stream = io.BytesIO()
        book.save(stream)
        return stream.getvalue()
